"""Validator: validates individual slides for content issues."""

from typing import Any

from pptx.slide import Slide

from ..utils import get_logger

logger = get_logger(__name__)


def validate_page(slide: Slide, slide_index: int) -> list[dict[str, Any]]:
    """
    Validate a single slide for common issues.

    Returns list of issues found (empty if valid).
    """
    issues = []

    for shape in slide.shapes:
        if not shape.is_placeholder or not shape.has_text_frame:
            continue

        ph_idx = shape.placeholder_format.idx if shape.placeholder_format else 0
        text = shape.text_frame.text

        # Check for empty required placeholders
        if not text.strip():
            ph_type = shape.placeholder_format.type if shape.placeholder_format else 0
            if ph_type in (1, 3):  # Title
                issues.append({
                    "slide": slide_index,
                    "placeholder": ph_idx,
                    "type": "empty_title",
                    "severity": "warning",
                    "message": "Title placeholder is empty",
                })

        # Check for excessively long text
        if len(text) > 800:
            issues.append({
                "slide": slide_index,
                "placeholder": ph_idx,
                "type": "text_too_long",
                "severity": "warning",
                "message": f"Text may overflow ({len(text)} characters)",
            })

        # Check for missing text frame wrapping
        try:
            tf = shape.text_frame
            if not tf.word_wrap:
                issues.append({
                    "slide": slide_index,
                    "placeholder": ph_idx,
                    "type": "word_wrap_disabled",
                    "severity": "info",
                    "message": "Word wrap is disabled on this placeholder",
                })
        except Exception:
            pass

    return issues
