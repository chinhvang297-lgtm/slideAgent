"""Self-corrector: detects and fixes common generation issues."""

from dataclasses import dataclass, field
from typing import Any

from pptx.presentation import Presentation
from pptx.util import Pt

from ..utils import get_logger

logger = get_logger(__name__)

MAX_TITLE_CHARS = 100
MAX_BODY_CHARS = 600
MAX_BULLET_CHARS = 120
MAX_BULLETS_PER_SLIDE = 8


@dataclass
class ValidationIssue:
    slide_index: int
    placeholder_idx: int
    issue_type: str
    description: str
    auto_fixed: bool = False
    fix_description: str = ""


@dataclass
class ValidationReport:
    slide_count: int = 0
    issues: list[ValidationIssue] = field(default_factory=list)

    @property
    def has_errors(self) -> bool:
        return any(not i.auto_fixed for i in self.issues)

    @property
    def error_count(self) -> int:
        return sum(1 for i in self.issues if not i.auto_fixed)

    @property
    def warning_count(self) -> int:
        return sum(1 for i in self.issues if i.auto_fixed)

    def summary(self) -> str:
        total = len(self.issues)
        if total == 0:
            return f"✓ {self.slide_count} slides validated with no issues"
        fixed = sum(1 for i in self.issues if i.auto_fixed)
        return (
            f"Validated {self.slide_count} slides: "
            f"{total} issues ({fixed} auto-fixed, {total-fixed} remaining)"
        )


class SelfCorrector:
    """Detects and auto-fixes common issues in the generated presentation."""

    def __init__(self) -> None:
        self.report = ValidationReport()

    def validate_and_fix_plan(
        self,
        slide_plan: dict[str, Any],
        template_schema: dict[str, Any],
    ) -> dict[str, Any]:
        """Validate and auto-fix a slide plan before generation."""
        slides = slide_plan.get("slides", [])
        self.report.slide_count = len(slides)

        fixed_slides = []
        for slide in slides:
            fixed_slide = self._fix_slide_plan(slide, template_schema)
            fixed_slides.append(fixed_slide)

        logger.info(self.report.summary())
        return {**slide_plan, "slides": fixed_slides}

    def validate_presentation(self, prs: Presentation) -> ValidationReport:
        """Validate a generated presentation for common issues."""
        report = ValidationReport(slide_count=len(prs.slides))

        for slide_idx, slide in enumerate(prs.slides):
            for shape in slide.shapes:
                if not shape.is_placeholder or not shape.has_text_frame:
                    continue

                ph_idx = shape.placeholder_format.idx if shape.placeholder_format else 0
                text = shape.text_frame.text

                # Check text overflow
                issue = self._check_text_overflow(slide_idx, ph_idx, shape, text)
                if issue:
                    report.issues.append(issue)

        return report

    def _fix_slide_plan(
        self,
        slide: dict[str, Any],
        template_schema: dict[str, Any],
    ) -> dict[str, Any]:
        """Auto-fix issues in a single slide plan."""
        layout_groups = {g["group_id"]: g for g in template_schema.get("layout_groups", [])}
        group_id = slide.get("layout_group_id", 0)
        group = layout_groups.get(group_id, {})

        fixed_blocks = []
        for block in slide.get("content_blocks", []):
            fixed_block = self._fix_content_block(block, slide["slide_number"])
            fixed_blocks.append(fixed_block)

        return {**slide, "content_blocks": fixed_blocks}

    def _fix_content_block(
        self,
        block: dict[str, Any],
        slide_number: int,
    ) -> dict[str, Any]:
        """Fix a single content block."""
        content_type = block.get("content_type", "text")
        ph_idx = block.get("placeholder_idx", 0)

        if content_type == "title":
            text = block.get("text", "")
            if len(text) > MAX_TITLE_CHARS:
                fixed_text = text[:MAX_TITLE_CHARS].rsplit(" ", 1)[0] + "..."
                self.report.issues.append(ValidationIssue(
                    slide_index=slide_number - 1,
                    placeholder_idx=ph_idx,
                    issue_type="title_overflow",
                    description=f"Title too long ({len(text)} chars > {MAX_TITLE_CHARS})",
                    auto_fixed=True,
                    fix_description=f"Truncated to '{fixed_text}'",
                ))
                return {**block, "text": fixed_text}

        elif content_type == "text":
            text = block.get("text", "")
            if len(text) > MAX_BODY_CHARS:
                fixed_text = text[:MAX_BODY_CHARS].rsplit(". ", 1)[0] + "."
                self.report.issues.append(ValidationIssue(
                    slide_index=slide_number - 1,
                    placeholder_idx=ph_idx,
                    issue_type="text_overflow",
                    description=f"Body text too long ({len(text)} chars > {MAX_BODY_CHARS})",
                    auto_fixed=True,
                    fix_description="Truncated to sentence boundary",
                ))
                return {**block, "text": fixed_text}

        elif content_type == "bullet_list":
            bullets = block.get("bullet_points", [])
            fixed = False
            fixed_bullets = []

            for bullet in bullets[:MAX_BULLETS_PER_SLIDE]:
                if len(bullet) > MAX_BULLET_CHARS:
                    fixed_bullet = bullet[:MAX_BULLET_CHARS].rsplit(" ", 1)[0] + "..."
                    fixed_bullets.append(fixed_bullet)
                    fixed = True
                else:
                    fixed_bullets.append(bullet)

            if len(bullets) > MAX_BULLETS_PER_SLIDE:
                self.report.issues.append(ValidationIssue(
                    slide_index=slide_number - 1,
                    placeholder_idx=ph_idx,
                    issue_type="too_many_bullets",
                    description=f"Too many bullets ({len(bullets)} > {MAX_BULLETS_PER_SLIDE})",
                    auto_fixed=True,
                    fix_description=f"Kept first {MAX_BULLETS_PER_SLIDE} bullets",
                ))

            if fixed or len(bullets) > MAX_BULLETS_PER_SLIDE:
                return {**block, "bullet_points": fixed_bullets}

        return block

    def _check_text_overflow(
        self,
        slide_idx: int,
        ph_idx: int,
        shape,
        text: str,
    ) -> ValidationIssue | None:
        """Check if text overflows the placeholder bounds."""
        try:
            tf = shape.text_frame
            if not tf.auto_size:
                # Estimate if text will overflow
                total_chars = len(text)
                ph_type = shape.placeholder_format.type if shape.placeholder_format else 0
                limit = MAX_TITLE_CHARS if ph_type in (1, 3) else MAX_BODY_CHARS

                if total_chars > limit:
                    return ValidationIssue(
                        slide_index=slide_idx,
                        placeholder_idx=ph_idx,
                        issue_type="potential_overflow",
                        description=f"Text may overflow ({total_chars} chars)",
                        auto_fixed=False,
                    )
        except Exception:
            pass
        return None
